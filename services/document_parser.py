import os
import logging
from pathlib import Path
import docx

logger = logging.getLogger(__name__)

class ParsingError(Exception):
    """Custom exception for document parsing failures."""
    pass

def extract_text_from_pdf(filepath: str) -> str:
    """
    Extract text content from a PDF document.
    Attempts PyMuPDF (fitz) first, with a fallback to PyPDF.
    Detects scanned / image-only PDFs that contain no extractable text.
    """
    text = ""
    # Method 1: PyMuPDF (fitz)
    try:
        import fitz
        doc = fitz.open(filepath)
        for page in doc:
            page_text = page.get_text("text")
            if page_text:
                text += page_text + "\n"
        doc.close()
    except Exception as e:
        logger.warning(f"PyMuPDF failed on {filepath}: {e}. Trying PyPDF fallback...")
        text = ""

    # Method 2: PyPDF Fallback
    if not text.strip():
        try:
            from pypdf import PdfReader
            reader = PdfReader(filepath)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        except Exception as e:
            logger.error(f"PyPDF failed on {filepath}: {e}")
            raise ParsingError(f"Failed to parse PDF file: {str(e)}")

    extracted_text = text.strip()
    if not extracted_text or len(extracted_text) < 20:
        raise ParsingError(
            "Extracted text is empty or extremely short. "
            "This PDF appears to be a scanned document or image-only PDF, which requires OCR (Optical Character Recognition) to process."
        )

    return extracted_text

def extract_text_from_docx(filepath: str) -> str:
    """
    Extract text content from a Microsoft Word (.docx) document.
    Extracts text from both paragraphs and table cells.
    """
    try:
        doc = docx.Document(filepath)
        full_text = []

        # Extract text from paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())

        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))

        extracted_text = "\n".join(full_text).strip()
        if not extracted_text:
            raise ParsingError("The DOCX document is empty or contains no readable text.")

        return extracted_text
    except ParsingError:
        raise
    except Exception as e:
        logger.error(f"Error reading DOCX file {filepath}: {e}")
        raise ParsingError(f"Could not read DOCX document: {str(e)}")

def extract_resume_text(filepath: str) -> str:
    """
    Unified façade function to extract raw text from supported resume formats (.pdf, .docx).
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"File not found: {filepath}")

    ext = Path(filepath).suffix.lower()
    if ext == '.pdf':
        return extract_text_from_pdf(filepath)
    elif ext == '.docx':
        return extract_text_from_docx(filepath)
    else:
        raise ValueError(f"Unsupported file format: {ext}. Only PDF and DOCX files are supported.")
